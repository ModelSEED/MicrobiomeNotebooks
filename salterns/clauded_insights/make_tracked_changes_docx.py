#!/usr/bin/env python
"""
Produce a copy of the manuscript DOCX with REAL Word tracked changes (w:ins / w:del)
implementing the review edits:
  * 5 reworded metabolic-modelling paragraphs (numbers corrected; text taken verbatim
    from manuscript_corrected.md) -> old paragraph content marked deleted, corrected
    text marked inserted;
  * new inserted paragraphs (all wrapped as insertions): geochem/gradient result,
    specificity caveat, modelling-as-attribution reframe, and a proposed-Figure-1 note;
  * (added 2026-10-26, dated DATE_R232FIX) the draft errors found in the R232 review:
    P35 "dereplicated medium and high quality MAGs" (35 of 310 exceed 10% contamination) and
    Table 3 (caption unbinned 68/89/89% -> 52/88/89%; values replaced cell by cell with the
    corrected 2020-family values from "Table 3 - corrected values (2020 GTDB families).csv").
    The third, P69 "(Figure S8)" -> "(Figure 7a)", is already covered by the P69 rewrite above.
Open in Word with Review > All Markup to see/accept the changes.
"""
import copy
import csv
from pathlib import Path
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(__file__).resolve().parents[1] / "Metabolic Manuscript.docx"
CORR = Path(__file__).resolve().parent / "manuscript_corrected.md"
OUT = Path(__file__).resolve().parent / "Metabolic Manuscript (tracked changes).docx"
TABLE3_CSV = Path(__file__).resolve().parents[1] / "Table 3 - corrected values (2020 GTDB families).csv"
AUTHOR, DATE = "Claude (review)", "2026-07-16T00:00:00Z"
DATE_R232FIX = "2026-10-26T00:00:00Z"   # the R232-review draft-error fixes (same author)
_id = iter(range(1000, 9000))


def _mkrun(text, italic=False, rpr=None):
    r = OxmlElement("w:r")
    if rpr is not None:                        # copy the formatting of the text being replaced
        r.append(copy.deepcopy(rpr))
    elif italic:
        rpr = OxmlElement("w:rPr"); rpr.append(OxmlElement("w:i")); r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; r.append(t)
    return r


def _wrap_ins(*runs, date=DATE):
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), str(next(_id))); ins.set(qn("w:author"), AUTHOR); ins.set(qn("w:date"), date)
    for r in runs:
        ins.append(r)
    return ins


def replace_paragraph(paragraph, new_text, date=DATE, keep_rpr=False, runs=None):
    """Mark ALL of the paragraph's content (runs + hyperlinked citations) deleted and
    insert new_text (tracked). Paragraphs here contain no pre-existing revisions.
    keep_rpr: give the inserted run the first deleted run's formatting (Table 3 cells);
    runs: insert these prepared w:r elements instead of new_text."""
    p = paragraph._p
    pPr = p.find(qn("w:pPr"))
    content = [c for c in list(p) if c is not pPr]
    first_r = p.find(qn("w:r"))
    rpr = first_r.find(qn("w:rPr")) if (keep_rpr and first_r is not None) else None
    if content:
        wdel = OxmlElement("w:del")
        wdel.set(qn("w:id"), str(next(_id))); wdel.set(qn("w:author"), AUTHOR); wdel.set(qn("w:date"), date)
        content[0].addprevious(wdel)
        for c in content:
            for t in c.iter(qn("w:t")):        # recurse into hyperlinks etc.
                t.tag = qn("w:delText")
            wdel.append(c)                     # move the whole child into the deletion
    p.append(_wrap_ins(*(runs if runs is not None else [_mkrun(new_text, rpr=rpr)]), date=date))


def _text_runs(p):
    return [r for r in p if r.tag == qn("w:r") and r.find(qn("w:t")) is not None]


def _split_run(r, k):
    """Split a single-w:t run at character k (both halves keep the run's formatting)."""
    right = copy.deepcopy(r)
    t, rt = r.find(qn("w:t")), right.find(qn("w:t"))
    t.text, rt.text = t.text[:k], t.text[k:]
    t.set(qn("xml:space"), "preserve"); rt.set(qn("xml:space"), "preserve")
    r.addnext(right)


def tracked_replace(paragraph, old, new, date):
    """Tracked replacement of the substring `old` (within the paragraph's direct text runs)
    by `new`; the inserted run copies the formatting of the first deleted run."""
    p = paragraph._p
    runs = _text_runs(p)
    assert all(len(r.findall(qn("w:t"))) == 1 for r in runs)
    full = "".join(r.find(qn("w:t")).text or "" for r in runs)
    assert full.count(old) == 1, f"tracked_replace anchor occurs {full.count(old)}x: {old!r}"
    start = full.index(old); end = start + len(old)
    for off in (start, end):                   # make run boundaries at start and end
        pos = 0
        for r in _text_runs(p):
            n = len(r.find(qn("w:t")).text or "")
            if pos < off < pos + n:
                _split_run(r, off - pos)
                break
            pos += n
    covered, pos = [], 0
    for r in _text_runs(p):
        n = len(r.find(qn("w:t")).text or "")
        if start <= pos and pos + n <= end and n:
            covered.append(r)
        pos += n
    assert "".join(r.find(qn("w:t")).text for r in covered) == old
    rpr = covered[0].find(qn("w:rPr"))
    wdel = OxmlElement("w:del")
    wdel.set(qn("w:id"), str(next(_id))); wdel.set(qn("w:author"), AUTHOR); wdel.set(qn("w:date"), date)
    covered[0].addprevious(wdel)
    for r in covered:
        for t in r.iter(qn("w:t")):
            t.tag = qn("w:delText")
        wdel.append(r)
    wdel.addnext(_wrap_ins(_mkrun(new, rpr=rpr), date=date))


def insert_after(paragraph, text, italic=False):
    """Insert a fully-tracked new paragraph after `paragraph`; returns the new Paragraph."""
    new_p = OxmlElement("w:p")
    src = paragraph._p.find(qn("w:pPr"))
    pPr = copy.deepcopy(src) if src is not None else OxmlElement("w:pPr")
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr"); pPr.append(rPr)
    ins_mark = OxmlElement("w:ins")
    ins_mark.set(qn("w:id"), str(next(_id))); ins_mark.set(qn("w:author"), AUTHOR); ins_mark.set(qn("w:date"), DATE)
    rPr.insert(0, ins_mark)                    # mark the paragraph-end as inserted
    new_p.append(pPr)
    new_p.append(_wrap_ins(_mkrun(text, italic)))
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


# ---- corrected paragraph text (verbatim from manuscript_corrected.md) --------------
corr_lines = [ln for ln in CORR.read_text().split("\n") if ln.strip()]
def corrected(anchor):
    hits = [ln for ln in corr_lines if anchor in ln]
    assert len(hits) == 1, f"corrected md anchor not unique ({len(hits)}): {anchor!r}"
    return hits[0].strip()

REPLACE = [   # anchor present in BOTH the original DOCX paragraph and its corrected md line
    "Soil DNA density was significantly positively correlated",
    "Roseovarius has been experimentally observed to consume",
    "betaine-consuming community members",
    "methylphosphonate-consuming organisms was significantly",
    "MPn-consuming organisms were",
]

# ---- new inserted paragraphs (a, c, b, Fig1) --------------------------------------
GEOCHEM = ("We further asked whether the environmental gradient itself, rather than any "
    "specific taxon, structures these emissions. Across the twelve cores, sediment sulfate "
    "(Spearman rho = 0.97, p < 0.001), chloride and salinity (rho = 0.94 and 0.79), and the "
    "inorganic N:P ratio (rho = 0.73, p = 0.007) were each more strongly correlated with "
    "methane emission than any biological variable, and a rank-based variance partition "
    "attributed only about 2% of the methane variance uniquely to methylphosphonate-consumer "
    "abundance beyond the sulfate gradient (42% unique to the gradient, 51% shared). This is "
    "consistent with a system organised by a single salinity/sulfate axis along which methane, "
    "metabolites, and community composition co-vary; the elevated N:P ratios (severe phosphate "
    "limitation) are themselves consistent with the methylphosphonate-scavenging hypothesis "
    "(proposed geochemistry figure).")
SPECIFICITY = ("Because the summed abundance of even a random set of MAGs of the same size "
    "correlates with methane about as well as the methylphosphonate- or betaine-consumer sets "
    "(37% and 45% of random subsets, respectively, matched or exceeded the observed "
    "coefficients), and because the association weakens after controlling for total DNA density, "
    "these cross-sample abundance correlations should be read as consistent with, rather than as "
    "specific evidence for, a particular methanogenic guild. The manipulative substrate-addition "
    "and culturing experiments provide the organism-specificity that correlation across the "
    "salinity gradient cannot.")
ATTRIBUTION = ("We therefore interpret the metabolic modelling as a demonstration of mechanistic "
    "sufficiency and taxonomic attribution -- methylphosphonate consumption can generate methane "
    "of the observed magnitude, and this capacity is concentrated in a few Rhodobacteraceae "
    "(Roseovarius, 77%; Albimonas, 11%; proposed Figure 7c) -- rather than as an independent "
    "statistical validation, which the six-core sample and the shared environmental gradient "
    "cannot support.")
FIG1 = ("[Proposed Figure 1: a conceptual overview of the study system -- the restoration / "
    "salinity-sulfate gradient, the two candidate methane pathways (aerobic methylphosphonate "
    "degradation and anaerobic betaine->TMA methylotrophic methanogenesis), and the five "
    "converging lines of evidence -- would orient the reader; a draft is provided in the "
    "clauded_insights folder.]")

doc = Document(str(SRC))

def find(anchor):
    hits = [p for p in doc.paragraphs if anchor in p.text]
    assert len(hits) == 1, f"DOCX anchor not unique ({len(hits)}): {anchor!r}"
    return hits[0]

# resolve ALL anchor paragraphs first (paragraph.text goes blank once wrapped in ins/del)
P = {a: find(a) for a in REPLACE}
P["_fig1"] = find("phnGHILMJ genes for the carbon-phosphorus lyase")
P["_p35"] = find("dereplicated medium and high quality metagenome-assembled genomes (MAGs)")
P["_t3cap"] = find("Table 3. The percent of total biomass")
T3 = doc.tables[2]
assert [c.text for c in T3.rows[0].cells] == ["GTDB-Tk Family", "Unrestored total", "Restored total",
                                               " Reference total"]

# 1) replace the 5 modelling paragraphs
for anchor in REPLACE:
    replace_paragraph(P[anchor], corrected(anchor))

# 2) insertions (using the stored paragraph objects)
a = insert_after(P["methylphosphonate-consuming organisms was significantly"], GEOCHEM)
insert_after(a, SPECIFICITY)
insert_after(P["MPn-consuming organisms were"], ATTRIBUTION)
insert_after(P["_fig1"], FIG1, italic=True)

# 3) R232-review draft errors (dated DATE_R232FIX; P69's "(Figure S8)" is gone with the P69 rewrite)
#    P35: 275 of the 310 MAGs meet MIMAG medium quality (35 exceed 10% CheckM contamination).
#    (The R232 review adds "Table S5" here; this copy has no Table S5.)
tracked_replace(P["_p35"], "dereplicated medium and high quality metagenome-assembled genomes (MAGs)",
                "310 dereplicated metagenome-assembled genomes (MAGs; 275 meet the MIMAG medium-quality criteria)",
                DATE_R232FIX)

#    Table 3: corrected 2020-family values (all 310 MAGs incl. archaea; R1+R2 / SF2 / R2A).  The old
#    table pooled R2A into "unrestored", dropped the archaea, and swapped Restored/Reference.
with open(TABLE3_CSV, newline="") as fh:
    t3csv = list(csv.DictReader(fh))
NEW_COLS = ("Unrestored (R1, R2)", "Restored (SF2)", "Reference (R2A)")
OLD_COLS = ("Old 'Unrestored total' (R1+R2+R2A)", "Old 'Restored total' (actually R2A)",
            "Old 'Reference total' (actually SF2)")
old_vals = {r["Family (2020 GTDB)"]: [r[c] for c in OLD_COLS] for r in t3csv if r[OLD_COLS[0]]}
top10 = [r for r in t3csv if r["In corrected top 10"] == "yes"]
unb = next(r for r in t3csv if r["Family (2020 GTDB)"].startswith("Unbinned share"))
assert len(top10) == 10 == len(T3.rows) - 1
tracked_replace(P["_t3cap"], "68%, 89%",
                ", ".join(f"{float(unb[c]) * 100:.0f}%" for c in NEW_COLS[:2]), DATE_R232FIX)
assert f"{float(unb[NEW_COLS[2]]) * 100:.0f}%" == "89%"     # the caption's third value is unchanged
# family-name formatting as in the old table: names italic, "undetermined" as it was, "(archaeal)" roman
_fam_rpr = {r.cells[0].text: copy.deepcopy(r.cells[0]._tc.find(".//" + qn("w:r")).find(qn("w:rPr")))
            for r in T3.rows[1:]}
_ital = _fam_rpr["Rhodobacteraceae"]
_roman = copy.deepcopy(_ital)
for _e in _roman.findall(qn("w:i")) + _roman.findall(qn("w:iCs")):
    _roman.remove(_e)


def family_runs(name):
    if name == "undetermined":
        return [_mkrun(name, rpr=_fam_rpr["undetermined"])]
    if name.endswith(" (archaeal)"):
        return [_mkrun(name[:-len(" (archaeal)")], rpr=_ital), _mkrun(" (archaeal)", rpr=_roman)]
    return [_mkrun(name, rpr=_ital)]


n_cells = 0
for row, new in zip(T3.rows[1:], top10):
    cells = row.cells
    fam_old = cells[0].text
    for c, ov in zip(cells[1:], old_vals[fam_old]):   # the docx table is the CSV's "old" table
        assert abs(float(c.text) - float(ov)) < 5e-7, (fam_old, c.text, ov)
    for ci, (c, nv) in enumerate(zip(cells, [new["Family (2020 GTDB)"],
                                             *(f"{float(new[k]):.6f}" for k in NEW_COLS)])):
        if c.text != nv:
            assert len(c.paragraphs) == 1
            replace_paragraph(c.paragraphs[0], nv, date=DATE_R232FIX, keep_rpr=True,
                              runs=family_runs(nv) if ci == 0 else None)
            n_cells += 1

doc.save(str(OUT))
print("saved", OUT.name)
print("tracked edits: 5 paragraph rewrites + 4 inserted paragraphs")
print(f"R232-review fixes: P35 phrase, Table 3 caption, {n_cells} Table 3 cells")
